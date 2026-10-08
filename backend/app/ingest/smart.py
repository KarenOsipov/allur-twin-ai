from __future__ import annotations

import csv
import difflib
import io
import json
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta

from app.domain.plant import Plant
from app.ingest.tables import UnsupportedFile, read_text

MAX_ROWS = 20_000
HEADER_SCAN = 15
SAMPLES = 4

KIND_TITLES = {
    "production": "Выпуск и план по сменам и линиям",
    "quality": "Брак и качество",
    "downtime": "Простои оборудования",
    "model_plan": "План по моделям",
    "staff": "Сотрудники",
}
IMPORTABLE = ("production", "quality", "downtime", "model_plan")


@dataclass(frozen=True)
class FieldSpec:
    name: str
    label: str
    type: str
    required: bool
    synonyms: tuple[str, ...]
    sniff: bool = False


def _f(name, label, type_, required, synonyms, sniff=False) -> FieldSpec:
    return FieldSpec(name, label, type_, required, tuple(synonyms), sniff)


_DAY = _f(
    "day",
    "Дата",
    "date",
    True,
    ["дата", "день", "число", "дата смены", "дата выпуска", "отчетная дата", "date", "day", "дт"],
    True,
)
_SHIFT = _f("shift", "Смена", "shift", False, ["смена", "номер смены", "№ смены", "shift", "shift no", "смен"], True)
_AREA = _f(
    "area",
    "Участок / линия",
    "area",
    True,
    [
        "участок",
        "линия",
        "цех",
        "производственная линия",
        "подразделение",
        "передел",
        "line",
        "area",
        "workshop",
        "station",
        "участок линия",
    ],
    True,
)

KINDS: dict[str, tuple[FieldSpec, ...]] = {
    "production": (
        _DAY,
        _SHIFT,
        _AREA,
        _f(
            "plan",
            "План, шт",
            "int",
            True,
            ["план", "плановый выпуск", "задание", "план шт", "план смены", "target", "plan", "planned", "норма"],
        ),
        _f(
            "fact",
            "Факт, шт",
            "int",
            True,
            [
                "факт",
                "фактический выпуск",
                "выпуск факт",
                "выпуск",
                "выпущено",
                "сделано",
                "выполнено",
                "actual",
                "output",
                "fact",
            ],
        ),
        _f(
            "run_hours",
            "Время работы, ч",
            "hours",
            False,
            [
                "время работы",
                "часы работы",
                "отработано",
                "работа ч",
                "фонд времени",
                "run hours",
                "runtime",
                "hours",
                "время",
            ],
        ),
        _f(
            "load_pct",
            "Загрузка, %",
            "pct",
            False,
            ["загрузка", "загрузка %", "коэффициент загрузки", "utilization", "load", "load %"],
        ),
    ),
    "quality": (
        _DAY,
        _SHIFT,
        _AREA,
        _f(
            "produced",
            "Выпущено, шт",
            "int",
            True,
            [
                "выпущено",
                "произведено",
                "проверено",
                "изготовлено",
                "всего",
                "produced",
                "inspected",
                "total",
                "выпуск",
            ],
        ),
        _f(
            "defects",
            "Брак, шт",
            "int",
            False,
            [
                "брак",
                "брак шт",
                "дефекты",
                "кол-во брака",
                "забраковано",
                "несоответствия",
                "defects",
                "scrap",
                "rejects",
                "defect",
            ],
        ),
        _f(
            "defect_pct",
            "Брак, %",
            "pct",
            False,
            [
                "% брака",
                "процент брака",
                "доля брака",
                "брак %",
                "уровень брака",
                "defect rate",
                "defect %",
                "scrap rate",
            ],
        ),
    ),
    "downtime": (
        _DAY,
        _SHIFT,
        _f("area", "Участок", "area", False, _AREA.synonyms, True),
        _f(
            "equipment",
            "Оборудование",
            "equipment",
            True,
            [
                "оборудование",
                "станок",
                "агрегат",
                "установка",
                "машина",
                "единица оборудования",
                "equipment",
                "machine",
                "asset",
                "робот",
            ],
            True,
        ),
        _f(
            "reason",
            "Причина",
            "text",
            False,
            [
                "причина",
                "причина простоя",
                "неисправность",
                "описание",
                "вид отказа",
                "reason",
                "cause",
                "failure",
                "комментарий",
            ],
        ),
        _f(
            "minutes",
            "Длительность, мин",
            "minutes",
            True,
            [
                "длительность",
                "продолжительность",
                "простой",
                "время простоя",
                "минут",
                "мин",
                "длит",
                "duration",
                "minutes",
                "downtime",
            ],
        ),
        _f(
            "started_at",
            "Начало",
            "time",
            False,
            ["начало", "время начала", "начало простоя", "с", "start", "started", "started at", "from"],
        ),
        _f(
            "planned",
            "Плановый",
            "bool",
            False,
            ["плановый", "плановый простой", "тип простоя", "план внеплан", "planned", "type"],
        ),
    ),
    "model_plan": (
        _f("month", "Месяц", "month", False, ["месяц", "период", "month", "period"], True),
        _f(
            "model",
            "Модель",
            "model",
            True,
            ["модель", "марка", "автомобиль", "модель авто", "наименование модели", "model", "car"],
            True,
        ),
        _f(
            "plan",
            "План, шт",
            "int",
            True,
            ["план", "план на месяц", "план выпуска", "количество", "кол-во", "plan", "qty", "quantity"],
        ),
    ),
    "staff": (
        _f(
            "name",
            "ФИО",
            "text",
            True,
            ["фио", "сотрудник", "имя", "фамилия имя", "работник", "name", "employee", "full name"],
        ),
        _f("position", "Должность", "text", False, ["должность", "позиция", "профессия", "position", "title", "role"]),
        _f("login", "Логин", "text", False, ["логин", "login", "user", "username"]),
        _f("area", "Участок", "area", False, _AREA.synonyms),
    ),
}
ONE_OF = {"quality": ("defects", "defect_pct")}
NUMERIC = {"int", "hours", "pct", "minutes"}
TITLE_HINTS = {
    "production": ("выпуск", "производ", "линии", "план-факт", "output", "production"),
    "quality": ("брак", "качеств", "дефект", "quality", "defect"),
    "downtime": ("простой", "простои", "оборудован", "отказ", "downtime"),
    "model_plan": ("план по модел", "модел", "производственный план", "model"),
    "staff": ("сотрудник", "персонал", "штат", "staff", "employee"),
}
AREA_ALIASES = {
    "WELD": ("свар", "weld", "кузов"),
    "PAINT": ("окрас", "покрас", "лкп", "paint", "краск"),
    "ASSY": ("сборк", "сбор", "assy", "assembl", "конвейер сборки"),
    "QC": ("контрол", "отк", "qc", "quality", "качеств", "инспекц"),
    "WH_IN": ("склад комплект", "склад компл", "комплект", "wh_in"),
    "WH_OUT": ("склад гп", "готов", "wh_out", "отгрузк"),
}
MONTHS = {
    "янв": 1,
    "фев": 2,
    "мар": 3,
    "апр": 4,
    "мая": 5,
    "май": 5,
    "июн": 6,
    "июл": 7,
    "авг": 8,
    "сен": 9,
    "окт": 10,
    "ноя": 11,
    "дек": 12,
    "jan": 1,
    "feb": 2,
    "mar": 3,
    "apr": 4,
    "may": 5,
    "jun": 6,
    "jul": 7,
    "aug": 8,
    "sep": 9,
    "oct": 10,
    "nov": 11,
    "dec": 12,
}
SHIFT_WORDS = {
    "i": 1,
    "ii": 2,
    "iii": 3,
    "перв": 1,
    "втор": 2,
    "трет": 3,
    "днев": 1,
    "утрен": 1,
    "вечер": 2,
    "ночн": 2,
    "day": 1,
    "night": 2,
    "a": 1,
    "b": 2,
    "а": 1,
    "б": 2,
}
TOTAL_WORDS = ("итого", "всего", "total", "сумма", "итог")
PLANNED_HINTS = ("плановое", "плановый", "замена", "то ", "обслуживание", "регламент")


@dataclass
class Sheet:
    name: str
    grid: list[list]
    lines: list[int] | None = None
    header_row: int = 0
    header: list[str] = field(default_factory=list)
    rows: list[tuple[int, list]] = field(default_factory=list)
    skipped_totals: int = 0


@dataclass
class Loaded:
    sheets: list[Sheet]
    fmt: str
    text: str = ""


def norm(s: object) -> str:
    t = str(s or "").lower().replace("ё", "е")
    t = re.sub(r"[_/\\|.,:;()\[\]{}\"'«»\-–—]+", " ", t)
    t = re.sub(r"[^\w%№ ]", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def _cell(v: object) -> object:
    if v is None:
        return None
    if isinstance(v, str):
        s = re.sub(r"\s+", " ", v).strip()
        return s or None
    if isinstance(v, bool):
        return v
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
        return int(v)
    return v


def load(filename: str, data: bytes) -> Loaded:
    name = filename.lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        _require_zip(data)
        return _load_xlsx(data)
    if name.endswith(".xls"):
        if data.startswith(b"PK"):
            return _load_xlsx(data)
        return _load_xls(data)
    if name.endswith((".csv", ".tsv", ".txt")):
        text, enc = decode(data)
        if name.endswith(".txt") and not _looks_delimited(text):
            doc = read_text(filename, text)
            return Loaded([_from_raw(t.title, t.header, t.rows) for t in doc.tables], f"текст · {enc}", doc.text)
        return _load_csv(filename, text, enc)
    if name.endswith(".json"):
        text, _ = decode(data)
        return load_json(filename, text)
    if name.endswith(".docx"):
        _require_zip(data)
        from app.ingest.tables import _read_docx

        doc = _read_docx(data)
        return Loaded([_from_raw(t.title, t.header, t.rows) for t in doc.tables], "Word (.docx)", doc.text)
    raise UnsupportedFile("Поддерживаются .xlsx, .xls, .csv, .tsv, .txt, .json и .docx")


def load_text(name: str, text: str) -> Loaded:
    if _looks_delimited(text):
        return _load_csv(name, text, "буфер обмена")
    doc = read_text(name, text)
    return Loaded([_from_raw(t.title, t.header, t.rows) for t in doc.tables], "текст", doc.text)


def _from_raw(title: str, header: list[str], rows: list[list[str]]) -> Sheet:
    return Sheet(title or "Таблица", [list(header), *[list(r) for r in rows]])


def _require_zip(data: bytes) -> None:
    if not data.startswith(b"PK") or not zipfile.is_zipfile(io.BytesIO(data)):
        raise UnsupportedFile("Файл повреждён или это не документ Office")


def decode(data: bytes) -> tuple[str, str]:
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8", errors="replace"), "UTF-8"
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16"), "UTF-16"
    try:
        return data.decode("utf-8"), "UTF-8"
    except UnicodeDecodeError:
        pass
    best, label, score = "", "", -1.0
    for enc, lab in (("cp1251", "Windows-1251"), ("cp866", "DOS-866"), ("koi8_r", "KOI8-R")):
        try:
            t = data.decode(enc)
        except UnicodeDecodeError:
            continue
        letters = sum(1 for ch in t if "а" <= ch <= "я")
        common = sum(1 for ch in t if ch in "оеаинтсрвл")
        bad = sum(1 for ch in t if ord(ch) > 0x2000 or ch in "ЎўЄєЇїЉЊЋЌЎЏ")
        sc = (common * 2 + letters - bad * 5) / max(len(t), 1)
        if sc > score:
            best, label, score = t, lab, sc
    if score < 0:
        raise UnsupportedFile("Не удалось определить кодировку (нужна UTF-8 или Windows-1251)")
    return best, label


def _looks_delimited(text: str) -> bool:
    lines = [ln for ln in text.splitlines()[:30] if ln.strip()]
    if len(lines) < 2:
        return False
    for d in (";", "\t", ","):
        counts = [ln.count(d) for ln in lines]
        if counts[0] >= 1 and sum(1 for c in counts if c == counts[0]) >= len(counts) * 0.7:
            return True
    return False


def _sniff_delimiter(text: str) -> str:
    lines = [ln for ln in text.splitlines()[:50] if ln.strip()]
    best, best_score = ";", -1.0
    for d in (";", "\t", ",", "|"):
        try:
            rows = list(csv.reader(io.StringIO("\n".join(lines)), delimiter=d))
        except csv.Error:
            continue
        widths = [len(r) for r in rows if r]
        if not widths:
            continue
        top = max(set(widths), key=widths.count)
        if top < 2:
            continue
        score = widths.count(top) / len(widths) * 10 + min(top, 12) * 0.1
        if score > best_score:
            best, best_score = d, score
    return best


def _load_csv(name: str, text: str, enc: str) -> Loaded:
    d = _sniff_delimiter(text)
    reader = csv.reader(io.StringIO(text), delimiter=d)
    numbered = [(n, [_cell(c) for c in row]) for n, row in enumerate(reader, 1)]
    numbered = [(n, r) for n, r in numbered if any(c is not None for c in r)][: MAX_ROWS + HEADER_SCAN]
    grid = [r for _, r in numbered]
    if len(grid) < 2:
        raise UnsupportedFile("В файле нет строк с данными")
    shown = {";": "«;»", ",": "«,»", "\t": "табуляция", "|": "«|»"}[d]
    title = re.sub(r"\.(csv|tsv|txt)$", "", name, flags=re.IGNORECASE)
    return Loaded([Sheet(title, grid, [n for n, _ in numbered])], f"CSV · {enc} · разделитель {shown}")


def _load_xlsx(data: bytes) -> Loaded:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets: list[Sheet] = []
    try:
        for ws in wb.worksheets:
            grid: list[list] = []
            lines: list[int] = []
            for n, row in enumerate(ws.iter_rows(values_only=True, max_row=MAX_ROWS + HEADER_SCAN), 1):
                values = [_cell(v) for v in row]
                if any(v is not None for v in values):
                    grid.append(values)
                    lines.append(n)
            if len(grid) >= 2:
                sheets.append(Sheet(ws.title, grid, lines))
    finally:
        wb.close()
    return Loaded(sheets, "Excel (.xlsx)")


def _load_xls(data: bytes) -> Loaded:
    try:
        import xlrd
    except ImportError as e:
        raise UnsupportedFile("Старый формат .xls: сохраните файл как .xlsx") from e
    try:
        book = xlrd.open_workbook(file_contents=data)
    except Exception as e:
        raise UnsupportedFile("Файл .xls повреждён или защищён паролем") from e
    sheets: list[Sheet] = []
    for sh in book.sheets():
        grid: list[list] = []
        lines: list[int] = []
        for r in range(min(sh.nrows, MAX_ROWS + HEADER_SCAN)):
            values = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        values.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                    except Exception:
                        values.append(cell.value)
                else:
                    values.append(_cell(cell.value if cell.value != "" else None))
            if any(v is not None for v in values):
                grid.append(values)
                lines.append(r + 1)
        if len(grid) >= 2:
            sheets.append(Sheet(sh.name, grid, lines))
    return Loaded(sheets, "Excel 97–2003 (.xls)")


def load_json(name: str, text: str) -> Loaded:
    try:
        obj = json.loads(text)
    except json.JSONDecodeError as e:
        raise UnsupportedFile(f"JSON с ошибкой: строка {e.lineno}, позиция {e.colno}") from e
    groups: list[tuple[str, object]] = []

    def walk(title: str, o: object, depth: int) -> None:
        if isinstance(o, list) and o:
            groups.append((title, o))
        elif isinstance(o, dict) and depth < 3:
            for k, v in o.items():
                walk(str(k), v, depth + 1)

    walk(re.sub(r"\.json$", "", name, flags=re.IGNORECASE), obj, 0)
    sheets: list[Sheet] = []
    for title, rows in groups:
        if all(isinstance(r, dict) for r in rows):
            header: list[str] = []
            for r in rows:
                for k in r:
                    if str(k) not in header:
                        header.append(str(k))
            grid = [header] + [[_cell(_json_value(r.get(k))) for k in header] for r in rows[:MAX_ROWS]]
        elif all(isinstance(r, list) for r in rows) and len(rows) >= 2:
            grid = [[_cell(_json_value(v)) for v in r] for r in rows[: MAX_ROWS + 1]]
        else:
            continue
        sheets.append(Sheet(title, grid))
    if not sheets:
        raise UnsupportedFile("В JSON нет таблиц: нужен список объектов или словарь «название: список объектов»")
    return Loaded(sheets, "JSON")


def _json_value(v: object) -> object:
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    return v


def synonym_score(header: str, synonyms: tuple[str, ...]) -> float:
    h = norm(header)
    if not h:
        return 0.0
    h_words = h.split()
    h_flat = h.replace(" ", "")
    best = 0.0
    for syn in synonyms:
        s = norm(syn)
        if not s:
            continue
        s_flat = s.replace(" ", "")
        if h == s or h_flat == s_flat:
            return 1.0
        if " " in s:
            if s in h:
                best = max(best, 0.92)
        elif len(s) <= 2:
            if s in h_words:
                best = max(best, 0.7)
        elif s in h_words:
            best = max(best, 0.9 if h_words[0] == s else 0.8)
        elif any(w.startswith(s) or (len(w) >= 4 and s.startswith(w)) for w in h_words):
            best = max(best, 0.8)
        elif len(s) >= 4 and s in h_flat:
            best = max(best, 0.7)
        else:
            r = difflib.SequenceMatcher(None, h_flat[: len(s_flat) + 3], s_flat).ratio()
            if r >= 0.8:
                best = max(best, r * 0.8)
    return best


def detect_header(sheet: Sheet) -> None:
    all_syn = tuple({s for spec in KINDS.values() for f in spec for s in f.synonyms if len(norm(s)) > 2})
    best_row, best_score = 0, -1.0
    for i, row in enumerate(sheet.grid[:HEADER_SCAN]):
        texts = [c for c in row if isinstance(c, str) and not _is_number(c) and _parse_date(c) is None]
        if len(texts) < 2:
            continue
        hits = sum(1 for c in texts if synonym_score(c, all_syn) >= 0.8)
        score = hits * 3 + len(texts) * 0.5 - i * 0.05
        if score > best_score:
            best_row, best_score = i, score
    sheet.header_row = best_row
    raw = sheet.grid[best_row]
    width = max(len(r) for r in sheet.grid[best_row : best_row + 200])
    header: list[str] = []
    for i in range(width):
        v = raw[i] if i < len(raw) else None
        header.append(str(v) if v is not None else f"Колонка {i + 1}")
    sheet.header = header
    sheet.rows = []
    lines = sheet.lines or list(range(1, len(sheet.grid) + 1))
    for k, row in enumerate(sheet.grid[best_row + 1 :], start=best_row + 1):
        n = lines[k]
        cells = list(row) + [None] * (width - len(row))
        if not any(c is not None for c in cells):
            continue
        first = next((c for c in cells if c is not None), None)
        if isinstance(first, str) and norm(first).split(" ")[0] in TOTAL_WORDS:
            sheet.skipped_totals += 1
            continue
        if [str(c) if c is not None else None for c in cells] == [str(h) for h in raw] + [None] * (width - len(raw)):
            continue
        sheet.rows.append((n, cells[:width]))
        if len(sheet.rows) >= MAX_ROWS:
            break


def _is_number(v: object) -> bool:
    if isinstance(v, bool):
        return False
    if isinstance(v, (int, float)):
        return True
    try:
        parse_number(v)
        return True
    except ValueError:
        return False


def parse_number(v: object) -> float:
    if isinstance(v, bool):
        raise ValueError("не число")
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v or "").strip().replace(" ", "").replace(" ", "").replace(" ", "").replace("'", "")
    s = re.sub(r"(шт|мин|ч|час|часов|%|₸|тг|kzt|ед)\.?$", "", s, flags=re.IGNORECASE)
    if not s:
        raise ValueError("пустое значение")
    if "," in s and "." in s:
        s = s.replace(",", "") if s.rfind(".") > s.rfind(",") else s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", ".")
    if not re.fullmatch(r"[-+]?\d+(\.\d+)?", s):
        raise ValueError(f"«{v}» — не число")
    return float(s)


def _parse_date(v: object) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if 30000 <= v <= 60000:
            return date(1899, 12, 30) + timedelta(days=int(v))
        return None
    if not isinstance(v, str):
        return None
    s = v.strip().lower().replace("г.", "").replace("г", "").strip()
    s = re.sub(r"\s+\d{1,2}:\d{2}(:\d{2})?$", "", s)
    for fmt in (
        "%d.%m.%Y",
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%d.%m.%y",
        "%d-%m-%Y",
        "%Y.%m.%d",
        "%d/%m/%y",
        "%Y-%m-%dT%H:%M:%S",
    ):
        try:
            d = datetime.strptime(s, fmt).date()
            return d if 2000 <= d.year <= 2100 else None
        except ValueError:
            continue
    m = re.fullmatch(r"(\d{1,2})\s+([а-яa-z]+)\.?\s+(\d{4})", s)
    if m:
        month = next((n for k, n in MONTHS.items() if m.group(2).startswith(k)), None)
        if month:
            try:
                return date(int(m.group(3)), month, int(m.group(1)))
            except ValueError:
                return None
    return None


def parse_month(v: object) -> str | None:
    if isinstance(v, (datetime, date)):
        return f"{v.year:04d}-{v.month:02d}"
    s = norm(v)
    m = re.fullmatch(r"(\d{4}) (\d{1,2})", s) or re.fullmatch(r"(\d{4})(\d{2})", s)
    if m and 1 <= int(m.group(2)) <= 12:
        return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}"
    m = re.fullmatch(r"(\d{1,2}) (\d{4})", s)
    if m and 1 <= int(m.group(1)) <= 12:
        return f"{int(m.group(2)):04d}-{int(m.group(1)):02d}"
    m = re.fullmatch(r"([а-яa-z]+) (\d{4})", s)
    if m:
        month = next((n for k, n in MONTHS.items() if m.group(1).startswith(k)), None)
        if month:
            return f"{int(m.group(2)):04d}-{month:02d}"
    d = _parse_date(v)
    return f"{d.year:04d}-{d.month:02d}" if d else None


def parse_shift(v: object) -> int | None:
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return int(v) if float(v).is_integer() and 1 <= v <= 4 else None
    s = norm(v)
    m = re.search(r"\b(\d)\b", s)
    if m and 1 <= int(m.group(1)) <= 4:
        return int(m.group(1))
    for w in s.replace("смена", "").replace("shift", "").split():
        for k, n in SHIFT_WORDS.items():
            if (len(k) <= 3 and w == k) or (len(k) > 3 and w.startswith(k)):
                return n
    return None


def parse_time(v: object, day: date | None) -> datetime | None:
    if isinstance(v, datetime):
        if v.year < 1910 and day:
            return datetime.combine(day, v.time())
        return v
    if isinstance(v, time):
        return datetime.combine(day, v) if day else None
    if isinstance(v, float) and 0 <= v < 1 and day:
        secs = round(v * 86400)
        return datetime.combine(day, time(secs // 3600 % 24, secs // 60 % 60))
    s = str(v or "").strip()
    m = re.fullmatch(r"(?:(\d{1,2})[./](\d{1,2})[./](\d{2,4})\s+)?(\d{1,2})[:.](\d{2})(?::\d{2})?", s)
    if m:
        h, mi = int(m.group(4)), int(m.group(5))
        if h > 23 or mi > 59:
            return None
        if m.group(1):
            y = int(m.group(3))
            base = date(y + 2000 if y < 100 else y, int(m.group(2)), int(m.group(1)))
        else:
            base = day
        return datetime.combine(base, time(h, mi)) if base else None
    try:
        return datetime.fromisoformat(s)
    except ValueError:
        return None


def parse_bool(v: object) -> bool | None:
    if isinstance(v, bool):
        return v
    s = norm(v)
    if not s:
        return None
    if s in ("да", "yes", "y", "true", "1", "+", "план", "плановый", "плановое", "planned"):
        return True
    if s in ("нет", "no", "n", "false", "0", "-", "внеплановый", "внеплан", "аварийный", "unplanned", "авария"):
        return False
    if s.startswith("внепл") or s.startswith("авар"):
        return False
    if s.startswith("план"):
        return True
    return None


class Resolver:
    def __init__(self, plant: Plant, shifts: set[int] | None = None) -> None:
        self.plant = plant
        self.shifts = shifts or {s.number for s in plant.shifts}
        self._eq = {self._key(e.code): e for e in plant.equipment}
        self._eq_names = {norm(e.name): e for e in plant.equipment}
        self._models = {}
        for m in plant.models:
            for k in (m.name, m.code, m.name.split()[-1]):
                self._models[norm(k)] = m.name

    @staticmethod
    def _key(s: object) -> str:
        return re.sub(r"[^0-9a-zа-я]", "", str(s).lower().replace("ё", "е"))

    def area(self, v: object) -> str | None:
        if v is None:
            return None
        direct = self.plant.area_by_name(str(v))
        if direct:
            return direct.code
        s = norm(v)
        flat = self._key(v)
        for a in self.plant.areas:
            if flat in (self._key(a.code), self._key(a.name), self._key(a.line or "")):
                return a.code
        s = re.sub(r"\b(цех|участок|линия|отдел|line|area|shop)\b", " ", s)
        s = re.sub(r"\b\d+\b", " ", s).strip()
        for code, stems in AREA_ALIASES.items():
            if any(st in s for st in stems) and any(a.code == code for a in self.plant.areas):
                return code
        for a in self.plant.areas:
            words = [w[:5] for w in norm(a.name).split() if len(w) >= 4]
            if words and all(w in s for w in words):
                return a.code
        return None

    def equipment(self, v: object) -> tuple[str, str | None]:
        raw = str(v).strip()
        e = self._eq.get(self._key(raw)) or self._eq_names.get(norm(raw))
        if e:
            return e.code, e.area
        for name, eq in self._eq_names.items():
            if norm(raw) and norm(raw) in name:
                return eq.code, eq.area
        return raw[:40], None

    def model(self, v: object) -> str:
        raw = str(v).strip()
        return self._models.get(norm(raw), raw[:60])

    def is_known_model(self, v: object) -> bool:
        return norm(v) in self._models

    def is_known_equipment(self, v: object) -> bool:
        return self.equipment(v)[1] is not None


@dataclass
class ColumnProfile:
    index: int
    name: str
    samples: list[str]
    filled: int
    ratios: dict[str, float]


def _sample_text(v: object) -> str:
    if isinstance(v, datetime):
        return v.strftime("%d.%m.%Y %H:%M") if (v.hour or v.minute) else v.strftime("%d.%m.%Y")
    if isinstance(v, date):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, time):
        return v.strftime("%H:%M")
    if isinstance(v, float):
        return f"{v:g}".replace(".", ",")
    return str(v)


def profile_columns(sheet: Sheet, res: Resolver) -> list[ColumnProfile]:
    out = []
    for i, name in enumerate(sheet.header):
        values = [r[i] for _, r in sheet.rows[:300] if i < len(r) and r[i] is not None]
        n = max(len(values), 1)

        def ratio(fn, vals=values, n=n) -> float:
            return sum(1 for v in vals if fn(v)) / n if vals else 0.0

        nums = [parse_number(v) for v in values if _is_number(v)]
        ratios = {
            "date": ratio(lambda v: _parse_date(v) is not None and not (isinstance(v, (int, float)) and v < 30000)),
            "num": ratio(_is_number),
            "int": ratio(lambda v: _is_number(v) and float(parse_number(v)).is_integer()),
            "shift": ratio(lambda v: parse_shift(v) in res.shifts) if (not nums or max(nums) <= 4) else 0.0,
            "area": ratio(lambda v: res.area(v) is not None),
            "equipment": ratio(res.is_known_equipment),
            "model": ratio(res.is_known_model),
            "text": ratio(lambda v: isinstance(v, str) and not _is_number(v) and _parse_date(v) is None),
            "month": ratio(lambda v: parse_month(v) is not None and not _is_number(v)),
            "time": ratio(lambda v: isinstance(v, (time, datetime)) or bool(re.search(r"\d{1,2}:\d{2}", str(v)))),
            "bool": ratio(lambda v: parse_bool(v) is not None),
        }
        uniq = []
        for v in values:
            t = _sample_text(v)
            if t not in uniq:
                uniq.append(t)
            if len(uniq) >= SAMPLES:
                break
        out.append(ColumnProfile(i, name, uniq, len(values), ratios))
    return out


def type_fit(spec: FieldSpec, col: ColumnProfile) -> float:
    r = col.ratios
    if col.filled == 0:
        return 0.3
    t = spec.type
    if t == "date":
        return r["date"]
    if t in NUMERIC:
        return r["num"]
    if t == "shift":
        return r["shift"]
    if t == "area":
        return max(r["area"], r["text"] * 0.5)
    if t == "equipment":
        return max(r["equipment"], r["text"] * 0.7)
    if t == "model":
        return max(r["model"], r["text"] * 0.7)
    if t == "month":
        return max(r["month"], r["date"])
    if t == "time":
        return r["time"]
    if t == "bool":
        return max(r["bool"], r["text"] * 0.5)
    return r["text"] if r["text"] else 0.3


def sniff_fit(spec: FieldSpec, col: ColumnProfile) -> float:
    r = col.ratios
    t = spec.type
    key = {
        "date": "date",
        "shift": "shift",
        "area": "area",
        "equipment": "equipment",
        "model": "model",
        "month": "month",
    }
    if t in key:
        return r[key[t]]
    return 0.0


def score_pair(spec: FieldSpec, col: ColumnProfile) -> tuple[float, str]:
    h = synonym_score(col.name, spec.synonyms)
    fit = type_fit(spec, col)
    if h >= 0.6:
        s = 0.75 * h + 0.25 * fit
        if fit < 0.5 and spec.type in NUMERIC | {"date", "shift"}:
            s *= 0.4 + fit
        return round(min(s, 1.0), 2), "по заголовку"
    if spec.sniff:
        sf = sniff_fit(spec, col)
        if sf >= 0.8:
            return round(0.4 + 0.25 * sf, 2), "по значениям"
    return 0.0, ""


@dataclass
class Assigned:
    column: int | None
    confidence: float
    how: str


def auto_map(kind: str, cols: list[ColumnProfile]) -> dict[str, Assigned]:
    specs = KINDS[kind]
    pairs = []
    for spec in specs:
        for col in cols:
            s, how = score_pair(spec, col)
            if s > 0:
                pairs.append((s, spec.required, spec.name, col.index, how))
    pairs.sort(key=lambda p: (-p[0], -p[1]))
    used_cols: set[int] = set()
    out: dict[str, Assigned] = {}
    for s, _req, name, ci, how in pairs:
        if name in out or ci in used_cols:
            continue
        out[name] = Assigned(ci, s, how)
        used_cols.add(ci)
    for spec in specs:
        out.setdefault(spec.name, Assigned(None, 0.0, ""))
    return out


def kind_confidence(kind: str, mapping: dict[str, Assigned], title: str) -> float:
    specs = KINDS[kind]
    req = [mapping[s.name].confidence for s in specs if s.required]
    if any(c <= 0 for c in req):
        return 0.0
    one_of = ONE_OF.get(kind)
    if one_of:
        best = max(mapping[n].confidence for n in one_of)
        if best <= 0:
            return 0.0
        req.append(best)
    conf = sum(req) / len(req)
    t = norm(title)
    if any(h in t for h in TITLE_HINTS[kind]):
        conf = min(1.0, conf + 0.08)
    return round(conf, 2)


def detect_kinds(cols: list[ColumnProfile], title: str) -> list[tuple[str, float, dict[str, Assigned]]]:
    found = []
    for kind in KINDS:
        m = auto_map(kind, cols)
        conf = kind_confidence(kind, m, title)
        if conf >= 0.45:
            found.append((kind, conf, m))
    names = {k for k, _, _ in found}
    if "production" in names or "quality" in names:
        found = [f for f in found if f[0] not in ("model_plan", "staff", "downtime") or f[1] > 0.9]
    if "downtime" in names:
        dt = next(f for f in found if f[0] == "downtime")
        found = [f for f in found if f[0] == "downtime" or f[1] > dt[1] + 0.1]
    if "staff" in names and len(found) > 1:
        found = [f for f in found if f[0] != "staff"]
    found.sort(key=lambda f: -f[1])
    return found


def unit_factor(kind: str, field_name: str, header: str) -> float:
    words = norm(header).split()
    if field_name == "minutes" and any(w in ("ч", "час", "часы", "часов", "hours", "h", "hrs") for w in words):
        return 60.0
    if field_name == "minutes" and any(w in ("сек", "секунд", "sec", "s") for w in words):
        return 1 / 60
    if field_name == "run_hours" and any(w.startswith("мин") or w == "min" for w in words):
        return 1 / 60
    return 1.0


def normalize_row(
    kind: str,
    mapping: dict[str, int | None],
    factors: dict[str, float],
    cells: list,
    res: Resolver,
    default_month: str,
) -> dict:
    def get(name: str) -> object:
        idx = mapping.get(name)
        if idx is None or idx >= len(cells):
            return None
        return cells[idx]

    rec: dict = {}
    if kind in ("production", "quality", "downtime"):
        raw = get("day")
        if raw is None:
            raise ValueError("нет даты")
        d = _parse_date(raw)
        if d is None:
            raise ValueError(f"дата «{_sample_text(raw)}» не распознана")
        rec["day"] = d
        raw = get("shift")
        if raw is None:
            rec["shift"] = 1
        else:
            sh = parse_shift(raw)
            if sh is None or sh not in res.shifts:
                raise ValueError(
                    f"смена «{_sample_text(raw)}» — на заводе смены {', '.join(map(str, sorted(res.shifts)))}"
                )
            rec["shift"] = sh

    def num(name: str, *, required: bool, lo: float = 0, hi: float = 1e9, integer: bool = False) -> float | None:
        raw = get(name)
        if raw is None:
            if required:
                raise ValueError(f"пусто: {_label(kind, name)}")
            return None
        try:
            v = parse_number(raw) * factors.get(name, 1.0)
        except ValueError as e:
            raise ValueError(f"{_label(kind, name)}: «{_sample_text(raw)}» — не число") from e
        if v < lo or v > hi:
            raise ValueError(f"{_label(kind, name)}: {v:g} вне допустимого диапазона {lo:g}–{hi:g}")
        return float(round(v)) if integer else round(v, 2)

    def pct(name: str) -> float | None:
        raw = get(name)
        if raw is None:
            return None
        try:
            v = parse_number(raw)
        except ValueError as e:
            raise ValueError(f"{_label(kind, name)}: «{_sample_text(raw)}» — не число") from e
        if isinstance(raw, (int, float)) and 0 < v <= 1 and not str(raw).endswith("%"):
            v *= 100
        if v < 0 or v > 150:
            raise ValueError(f"{_label(kind, name)}: {v:g}% вне диапазона")
        return round(v, 2)

    if kind in ("production", "quality"):
        raw = get("area")
        if raw is None:
            raise ValueError("не указан участок")
        code = res.area(raw)
        if code is None:
            raise ValueError(f"участок «{_sample_text(raw)}» не найден в схеме завода")
        if kind == "production" and not res.plant.area(code).line:
            raise ValueError(f"«{_sample_text(raw)}» — не производственная линия (нужны сварка, окраска или сборка)")
        rec["area"] = code
    if kind == "production":
        rec["plan"] = num("plan", required=True, hi=100_000, integer=True)
        rec["fact"] = num("fact", required=True, hi=100_000, integer=True)
        rh = num("run_hours", required=False, hi=24)
        if rh is not None:
            rec["run_hours"] = rh
        lp = pct("load_pct")
        if lp is not None:
            rec["load_pct"] = lp
    elif kind == "quality":
        produced = num("produced", required=True, hi=100_000, integer=True)
        defects = num("defects", required=False, hi=100_000, integer=True)
        if defects is None:
            p = pct("defect_pct")
            if p is None:
                raise ValueError("нет ни количества брака, ни процента брака")
            defects = float(round(produced * p / 100))
        if defects > produced:
            raise ValueError(f"брак {defects:g} больше выпуска {produced:g}")
        rec["produced"] = produced
        rec["defects"] = defects
    elif kind == "downtime":
        raw = get("equipment")
        if raw is None:
            raise ValueError("не указано оборудование")
        code, eq_area = res.equipment(raw)
        rec["equipment"] = code
        area_raw = get("area")
        area = res.area(area_raw) if area_raw is not None else None
        if area_raw is not None and area is None and eq_area is None:
            raise ValueError(f"участок «{_sample_text(area_raw)}» не найден в схеме завода")
        rec["area"] = area or eq_area
        if rec["area"] is None:
            raise ValueError(f"оборудование «{code}» неизвестно и участок не указан")
        reason = get("reason")
        rec["reason"] = str(reason).strip()[:120] if reason is not None else "Причина не указана"
        rec["minutes"] = num("minutes", required=True, hi=1440)
        if not rec["minutes"]:
            raise ValueError("длительность простоя 0 мин")
        st = get("started_at")
        if st is not None:
            t = parse_time(st, rec["day"])
            if t is None:
                raise ValueError(f"время начала «{_sample_text(st)}» не распознано")
            rec["started_at"] = t.replace(second=0, microsecond=0)
        pl = get("planned")
        flag = parse_bool(pl) if pl is not None else None
        if flag is None:
            flag = any(h in (rec["reason"].lower() + " ") for h in PLANNED_HINTS)
        rec["planned"] = flag
    elif kind == "model_plan":
        raw = get("model")
        if raw is None:
            raise ValueError("не указана модель")
        rec["model"] = res.model(raw)
        raw = get("month")
        month = parse_month(raw) if raw is not None else default_month
        if month is None:
            raise ValueError(f"месяц «{_sample_text(raw)}» не распознан")
        rec["month"] = month
        rec["plan"] = num("plan", required=True, hi=1_000_000, integer=True)
    elif kind == "staff":
        raw = get("name")
        if raw is None:
            raise ValueError("нет ФИО")
        rec["name"] = str(raw)
        for k in ("position", "login"):
            v = get(k)
            if v is not None:
                rec[k] = str(v)
    return rec


def _label(kind: str, name: str) -> str:
    for f in KINDS[kind]:
        if f.name == name:
            return f.label
    return name


def natural_key(kind: str, rec: dict) -> tuple:
    if kind in ("production", "quality"):
        return (rec["day"], rec["shift"], rec["area"])
    if kind == "downtime":
        if rec.get("started_at"):
            return (rec["day"], rec["shift"], rec["equipment"], rec["started_at"])
        return (rec["day"], rec["shift"], rec["equipment"], rec["reason"].lower())
    if kind == "model_plan":
        return (rec["month"], rec["model"])
    return (rec.get("name"),)


COMPARE = {
    "production": ("plan", "fact", "run_hours", "load_pct"),
    "quality": ("produced", "defects"),
    "downtime": ("area", "reason", "minutes", "planned"),
    "model_plan": ("plan",),
}
FIELD_LABELS = {
    "plan": "План",
    "fact": "Факт",
    "run_hours": "Время работы, ч",
    "load_pct": "Загрузка, %",
    "produced": "Выпущено",
    "defects": "Брак",
    "area": "Участок",
    "reason": "Причина",
    "minutes": "Минуты",
    "planned": "Плановый",
}
