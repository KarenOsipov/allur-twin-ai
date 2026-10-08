from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BRAND = "E3241B"
INK = "1D1515"
MUTED = "7D6E6C"


@dataclass
class Sheet:
    title: str
    columns: list[str]
    rows: list[list] = field(default_factory=list)
    note: str | None = None


def workbook(report: str, meta: list[tuple[str, str]], sheets: list[Sheet]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for sh in sheets:
        ws = wb.create_sheet(_safe_title(sh.title))
        ws.sheet_view.showGridLines = False
        ws["A1"] = "allur · АО «Группа компаний АЛЛЮР» · Сборочный завод, г. Костанай"
        ws["A1"].font = Font(bold=True, size=13, color=BRAND)
        ws["A2"] = f"{report} — {sh.title}"
        ws["A2"].font = Font(bold=True, size=12, color=INK)
        r = 3
        for k, v in meta:
            ws.cell(r, 1, k).font = Font(color=MUTED, size=10)
            ws.cell(r, 2, v).font = Font(color=INK, size=10)
            r += 1
        if sh.note:
            ws.cell(r, 1, sh.note).font = Font(italic=True, color=MUTED, size=10)
            r += 1
        r += 1
        header = r
        for i, col in enumerate(sh.columns, 1):
            c = ws.cell(header, i, col)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=BRAND)
            c.alignment = Alignment(vertical="center", wrap_text=True)
        widths = [len(c) for c in sh.columns]
        for row in sh.rows:
            r += 1
            for i, v in enumerate(row, 1):
                cell = ws.cell(r, i, _value(v))
                if isinstance(v, str):
                    cell.data_type = "s"
                    cell.alignment = Alignment(vertical="top", wrap_text=len(v) > 60)
                elif isinstance(v, datetime):
                    cell.number_format = "dd.mm.yyyy hh:mm"
                elif isinstance(v, date):
                    cell.number_format = "dd.mm.yyyy"
                elif isinstance(v, float):
                    cell.number_format = "#,##0.0" if abs(v) < 1000 else "#,##0"
                widths[i - 1] = max(widths[i - 1], min(len(_text(v)), 70))
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(w + 2, 72))
        ws.freeze_panes = ws.cell(header + 1, 1)
        if sh.rows:
            ws.auto_filter.ref = f"A{header}:{get_column_letter(len(sh.columns))}{r}"
    if not wb.sheetnames:
        wb.create_sheet("Пусто")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _value(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "да" if v else "нет"
    if isinstance(v, (int, float, datetime, date)):
        return v
    if isinstance(v, (list, tuple)):
        return ", ".join(_text(x) for x in v)
    if isinstance(v, dict):
        return "; ".join(f"{k}: {_text(x)}" for k, x in v.items())
    return str(v)


def _text(v) -> str:
    if isinstance(v, datetime):
        return v.strftime("%d.%m.%Y %H:%M")
    if isinstance(v, date):
        return v.strftime("%d.%m.%Y")
    return str(_value(v))


def _safe_title(t: str) -> str:
    for ch in "[]:*?/\\":
        t = t.replace(ch, " ")
    return t[:31] or "Лист"
