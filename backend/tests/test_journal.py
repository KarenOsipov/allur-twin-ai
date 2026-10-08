from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from tests.conftest import login, user_id
from tests.test_sandbox import _fork


def _xlsx(res):
    assert res.status_code == 200, res.text
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats")
    return load_workbook(io.BytesIO(res.content))


def test_journal_records_people_and_plant_events(client):
    bad = client.post("/api/v1/auth/login", json={"user_id": user_id(client, "supervisor"), "pin": "9999"})
    assert bad.status_code == 401
    h = login(client, "supervisor")
    client.app.state.container.live.sim = _fork()
    client.post("/api/v1/floor/failure", headers=h, json={"equipment": "ABB-01", "minutes": 20})
    client.post("/api/v1/assistant", headers=h, json={"question": "Где узкое место?"})

    j = client.get("/api/v1/journal", headers=h).json()
    actions = {(e["category"], e["action"]) for e in j["items"]}
    assert {("auth", "login_failed"), ("auth", "login"), ("control", "failure"), ("assistant", "ask")} <= actions
    fail = next(e for e in j["items"] if e["action"] == "failure")
    assert fail["actor"] == "Ерлан Жумабеков" and fail["severity"] == "warning"
    assert all("ip" not in e["details"] for e in j["items"])
    admin = login(client, "admin", "0000")
    assert any("ip" in e["details"] for e in client.get("/api/v1/journal?category=auth", headers=admin).json()["items"])

    only = client.get("/api/v1/journal?category=control", headers=h).json()
    assert only["total"] >= 1 and all(e["category"] == "control" for e in only["items"])
    found = client.get("/api/v1/journal", headers=h, params={"q": "ABB-01"}).json()
    assert found["total"] >= 1


def test_journal_export_excel_is_safe_from_formulas(client):
    h = login(client, "admin", "0000")
    client.post("/api/v1/assistant", headers=h, json={"question": '=HYPERLINK("http://x","клик")'})
    wb = _xlsx(client.get("/api/v1/journal/export.xlsx", headers=h))
    ws = wb["Журнал"]
    cells = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and "HYPERLINK" in c.value]
    assert cells and all(c.data_type == "s" for c in cells)
    csv = client.get("/api/v1/journal/export.csv", headers=h)
    assert csv.status_code == 200 and "Событие" in csv.text


@pytest.mark.parametrize("page", ["floor", "kpi", "quality", "forecast", "incidents", "data", "ai"])
def test_every_page_exports_to_excel(client, page):
    h = login(client, "director", "1111")
    wb = _xlsx(client.get(f"/api/v1/export/{page}.xlsx", headers=h, params={"note": "Проверить и доработать"}))
    assert wb.sheetnames
    first = wb[wb.sheetnames[0]]
    assert "allur" in str(first["A1"].value)
    assert any(c.value == "Проверить и доработать" for row in first.iter_rows(max_row=8) for c in row)


def test_floor_export_in_simulation_has_summary(client):
    h = login(client, "supervisor")
    client.app.state.container.live.sim = _fork()
    sb = client.post(
        "/api/v1/sandbox", headers=h, json={"events": [{"kind": "failure", "equipment": "Конвейер-03", "minutes": 40}]}
    )
    hs = {**h, "X-Sandbox": sb.json()["id"]}
    wb = _xlsx(client.get("/api/v1/export/floor.xlsx", headers=hs))
    assert "Итог симуляции" in wb.sheetnames and "Хронология" in wb.sheetnames
    journal = client.get("/api/v1/journal?category=simulation", headers=h).json()
    assert journal["total"] >= 1


def test_builder_layouts_crud_and_journal(client):
    h = login(client, "supervisor")
    data = {
        "nodes": [
            {"id": "a", "type": "source", "position": {"x": 0, "y": 0}, "data": {"kind": "source", "name": "Склад"}},
            {
                "id": "b",
                "type": "station",
                "position": {"x": 300, "y": 0},
                "data": {"kind": "station", "name": "Участок", "equipment": [{"code": "X-1"}, {"code": "X-2"}]},
            },
        ],
        "edges": [{"id": "e", "source": "a", "target": "b"}],
    }
    made = client.post("/api/v1/layouts", headers=h, json={"name": "Цех №2", "data": data}).json()
    assert made["nodes"] == 2 and made["equipment"] == 2
    assert client.get("/api/v1/layouts", headers=h).json()[0]["name"] == "Цех №2"
    got = client.get(f"/api/v1/layouts/{made['id']}", headers=h).json()
    assert got["data"]["edges"][0]["target"] == "b"
    client.put(f"/api/v1/layouts/{made['id']}", headers=h, json={"name": "Цех №3", "data": data})
    bad = client.post("/api/v1/layouts", headers=h, json={"name": "x", "data": {"nodes": "нет"}})
    assert bad.status_code == 422
    assert client.delete(f"/api/v1/layouts/{made['id']}", headers=h).status_code == 200
    j = client.get("/api/v1/journal?category=builder", headers=h).json()
    assert {e["action"] for e in j["items"]} >= {"save", "update", "delete"}
