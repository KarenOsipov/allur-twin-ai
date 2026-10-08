from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from enum import StrEnum

MAX_ROWS = 20_000


class TableKind(StrEnum):
    PRODUCTION = "production"
    DOWNTIME = "downtime"
    MODEL_PLAN = "model_plan"
    QUALITY = "quality"


TABLE_TITLES = {
    TableKind.PRODUCTION: "Работа производственных линий",
    TableKind.DOWNTIME: "Простои оборудования",
    TableKind.MODEL_PLAN: "Производственный план",
    TableKind.QUALITY: "Показатели качества",
}


@dataclass
class RawTable:
    title: str
    header: list[str]
    rows: list[list[str]]


@dataclass
class ParsedDocument:
    tables: list[RawTable]
    text: str = ""


@dataclass
class TableResult:
    kind: TableKind
    title: str
    records: list[dict]
    errors: list[str] = field(default_factory=list)


class UnsupportedFile(ValueError):
    pass


def read_document(filename: str, data: bytes) -> ParsedDocument:
    name = filename.lower()
    if name.endswith(".docx"):
        _require_zip(data)
        return _read_docx(data)
    if name.endswith(".xlsx"):
        _require_zip(data)
        return _read_xlsx(data)
    if name.endswith(".csv"):
        return _read_csv(name, data)
    if name.endswith(".json"):
        return read_json(name, _decode(data))
    if name.endswith((".txt", ".tsv", ".text", ".md")):
        return read_text(name, _decode(data))
    if name.endswith(".xls"):
        raise UnsupportedFile("Старый формат .xls: откройте файл в Excel и сохраните как .xlsx")
    raise UnsupportedFile("Поддерживаются .docx, .xlsx, .csv, .tsv, .txt и .json — или вставьте таблицу текстом")


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise UnsupportedFile("Не удалось определить кодировку (нужна UTF-8 или Windows-1251)")


def read_json(name: str, text: str) -> ParsedDocument:
    import json

    try:
        obj = json.loads(text)
    except json.JSONDecodeError as e:
        raise UnsupportedFile(f"JSON с ошибкой: строка {e.lineno}, позиция {e.colno}") from e
    groups = obj.items() if isinstance(obj, dict) else [(name, obj)]
    tables: list[RawTable] = []
    for title, rows in groups:
        if not isinstance(rows, list) or not rows or not all(isinstance(r, dict) for r in rows):
            continue
        header: list[str] = []
        for r in rows:
            for k in r:
                if k not in header:
                    header.append(str(k))
        grid = [[_cell("" if r.get(k) is None else str(r.get(k))) for k in header] for r in rows[:MAX_ROWS]]
        tables.append(RawTable(str(title), header, grid))
    if not tables:
        raise UnsupportedFile("В JSON нет таблиц: нужен список объектов или словарь «название: список объектов»")
    return ParsedDocument(tables)


def read_text(name: str, text: str) -> ParsedDocument:
    tables: list[RawTable] = []
    title = name
    block: list[str] = []

    def flush() -> None:
        nonlocal block, title
        lines = [ln for ln in block if ln.strip() and not re.fullmatch(r"[\s|:+-]+", ln)]
        block = []
        if len(lines) < 2:
            if len(lines) == 1:
                title = lines[0].strip()
            return
        sample = "\n".join(lines[:20])
        delim = max(("\t", ";", "|", ","), key=lambda d: (sum(1 for ln in lines[:20] if d in ln), sample.count(d)))
        while len(lines) > 2 and delim not in lines[0] and delim in lines[1]:
            title = lines.pop(0).strip()
        if not any(delim in ln for ln in lines):
            grid = [re.split(r"\s{2,}", ln.strip()) for ln in lines]
        elif delim == "|":
            grid = [[c for c in ln.strip().strip("|").split("|")] for ln in lines]
        else:
            grid = list(csv.reader(io.StringIO("\n".join(lines)), delimiter=delim))
        grid = [[_cell(c) for c in r] for r in grid if any(c.strip() for c in r)]
        if len(grid) >= 2:
            tables.append(RawTable(title, grid[0], grid[1 : MAX_ROWS + 1]))
        title = name

    for ln in text.splitlines():
        if not ln.strip():
            flush()
        else:
            block.append(ln)
    flush()
    if not tables:
        raise UnsupportedFile("Не нашли таблицу в тексте: нужна строка заголовков и строки данных")
    return ParsedDocument(tables, text)


def _require_zip(data: bytes) -> None:
    if not data.startswith(b"PK") or not zipfile.is_zipfile(io.BytesIO(data)):
        raise UnsupportedFile("Файл повреждён или это не документ Office")


def _cell(text: str | None) -> str:
    return re.sub(r"\s+", " ", (text or "")).strip()


def _read_docx(data: bytes) -> ParsedDocument:
    from docx import Document

    doc = Document(io.BytesIO(data))
    body = doc.element.body
    tables: list[RawTable] = []
    paragraphs: list[str] = []
    last_heading = ""
    t_index = 0
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = _cell("".join(node.text or "" for node in child.iter() if node.tag.endswith("}t")))
            if text:
                paragraphs.append(text)
                last_heading = text
        elif tag == "tbl":
            table = doc.tables[t_index]
            t_index += 1
            grid = [[_cell(c.text) for c in row.cells] for row in table.rows]
            grid = [r for r in grid if any(r)]
            if len(grid) >= 2:
                tables.append(RawTable(last_heading, grid[0], grid[1:MAX_ROWS]))
    return ParsedDocument(tables, "\n".join(paragraphs))


def _read_xlsx(data: bytes) -> ParsedDocument:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    tables: list[RawTable] = []
    for ws in wb.worksheets:
        grid: list[list[str]] = []
        for row in ws.iter_rows(values_only=True, max_row=MAX_ROWS):
            values = [_xlsx_value(v) for v in row]
            if any(values):
                grid.append(values)
        if len(grid) >= 2:
            tables.append(RawTable(ws.title, grid[0], grid[1:]))
    wb.close()
    return ParsedDocument(tables)


def _xlsx_value(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, date):
        return v.strftime("%d.%m.%Y")
    return _cell(str(v))


def _read_csv(name: str, data: bytes) -> ParsedDocument:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise UnsupportedFile("Не удалось определить кодировку CSV (нужна UTF-8 или Windows-1251)")
    sample = text[:4096]
    delimiter = ";" if sample.count(";") >= sample.count(",") else ","
    grid = [[_cell(c) for c in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    grid = [r for r in grid if any(r)][: MAX_ROWS + 1]
    if len(grid) < 2:
        raise UnsupportedFile("В CSV нет строк с данными")
    return ParsedDocument([RawTable(name, grid[0], grid[1:])])


SYNONYMS: list[tuple[str, str]] = [
    ("date", "дата"),
    ("day", "дата"),
    ("день", "дата"),
    ("число", "дата"),
    ("productionline", "линия"),
    ("line", "линия"),
    ("линияучасток", "линия"),
    ("plan", "план"),
    ("target", "план"),
    ("задание", "план"),
    ("actual", "факт"),
    ("fact", "факт"),
    ("output", "факт"),
    ("выпускфакт", "факт"),
    ("выпуск", "факт"),
    ("сделано", "факт"),
    ("equipment", "оборудование"),
    ("machine", "оборудование"),
    ("станок", "оборудование"),
    ("агрегат", "оборудование"),
    ("установка", "оборудование"),
    ("reason", "причина"),
    ("cause", "причина"),
    ("duration", "длительность"),
    ("minutes", "длительность"),
    ("downtime", "длительность"),
    ("минут", "длительность"),
    ("продолжительность", "длительность"),
    ("produced", "выпущено"),
    ("произведено", "выпущено"),
    ("изготовлено", "выпущено"),
    ("defects", "брак"),
    ("дефект", "брак"),
    ("defect", "брак"),
    ("scrap", "брак"),
    ("model", "модель"),
    ("марка", "модель"),
    ("area", "участок"),
    ("цех", "участок"),
    ("station", "участок"),
    ("shift", "смена"),
    ("load", "загрузка"),
    ("runhours", "времяработы"),
    ("hours", "времяработы"),
]


def _norm(h: str) -> str:
    n = re.sub(r"[^а-яёa-z%]", "", h.lower())
    for syn, canon in SYNONYMS:
        if n.startswith(syn):
            return canon + n[len(syn) :]
    return n


def classify(header: list[str]) -> TableKind | None:
    h = {_norm(x) for x in header}
    joined = " ".join(h)

    def has(prefix: str) -> bool:
        return any(x.startswith(prefix) for x in h)

    if has("линия") and has("план") and has("факт"):
        return TableKind.PRODUCTION
    if has("участок") and has("план") and has("факт") and not has("оборудование"):
        return TableKind.PRODUCTION
    if has("оборудование") and (has("причина") or has("длительность")):
        return TableKind.DOWNTIME
    if has("брак") and (has("выпущено") or has("факт")):
        return TableKind.QUALITY
    if has("модель") and "план" in joined:
        return TableKind.MODEL_PLAN
    return None


def _col(header: list[str], *names: str) -> int | None:
    normed = [_norm(x) for x in header]
    for name in names:
        for i, h in enumerate(normed):
            if h.startswith(name):
                return i
    return None


def parse_date(s: str) -> date:
    s = s.strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"дата «{s}» не распознана")


def parse_number(s: str) -> float:
    cleaned = s.replace(" ", "").replace(" ", "").replace(",", ".").rstrip("%")
    if not cleaned:
        raise ValueError("пустое значение")
    return float(cleaned)


def extract(table: RawTable) -> TableResult | None:
    kind = classify(table.header)
    if kind is None:
        return None
    h = table.header
    result = TableResult(kind, TABLE_TITLES[kind], [])
    cols = {
        TableKind.PRODUCTION: {
            "day": _col(h, "дата"),
            "line": _col(h, "линия", "участок"),
            "plan": _col(h, "план"),
            "fact": _col(h, "факт"),
            "run_hours": _col(h, "времяработы", "время"),
            "load_pct": _col(h, "загрузка"),
            "shift": _col(h, "смена"),
        },
        TableKind.DOWNTIME: {
            "day": _col(h, "дата"),
            "area": _col(h, "участок"),
            "equipment": _col(h, "оборудование"),
            "reason": _col(h, "причина"),
            "minutes": _col(h, "длительность", "простой"),
            "shift": _col(h, "смена"),
        },
        TableKind.MODEL_PLAN: {"model": _col(h, "модель"), "plan": _col(h, "план")},
        TableKind.QUALITY: {
            "day": _col(h, "дата"),
            "area": _col(h, "участок"),
            "produced": _col(h, "выпущено", "факт"),
            "defects": _col(h, "брак"),
            "shift": _col(h, "смена"),
        },
    }[kind]
    optional = ("shift", "load_pct", "run_hours") + (("reason",) if kind == TableKind.DOWNTIME else ())
    required = [k for k, v in cols.items() if k not in optional and v is None]
    if required:
        result.errors.append(f"не найдены колонки: {', '.join(required)}")
        return result

    numeric = {
        "plan",
        "fact",
        "run_hours",
        "load_pct",
        "minutes",
        "produced",
        "defects",
        "shift",
    }
    for n, row in enumerate(table.rows, start=2):
        rec: dict = {}
        try:
            for key, idx in cols.items():
                if idx is None or idx >= len(row):
                    continue
                raw = row[idx]
                if key == "day":
                    rec[key] = parse_date(raw)
                elif key in numeric:
                    if raw == "" and key in ("shift", "load_pct"):
                        continue
                    rec[key] = parse_number(raw)
                else:
                    if not raw:
                        raise ValueError(f"пустое поле «{key}»")
                    rec[key] = raw
            if kind == TableKind.DOWNTIME and not rec.get("reason"):
                rec["reason"] = "Причина не указана"
            result.records.append(rec)
        except ValueError as e:
            result.errors.append(f"строка {n}: {e}")
    return result


_TARGET_PATTERNS = {
    "oee_pct": r"OEE[^0-9]*(\d+[.,]?\d*)\s*%",
    "defect_pct": r"брак[а-я]*[^0-9]*(\d+[.,]?\d*)\s*%",
    "critical_downtime_min_per_day": r"просто[йя][^0-9]*(\d+)\s*мин",
    "month_output": r"план выпуска[^0-9]*([\d\s ]+)\s*автомоб",
    "shifts": r"(\d+)\s*смен",
    "shift_hours": r"смен[а-я]*\s*по\s*(\d+)\s*час",
}


def extract_targets(text: str) -> dict[str, float]:
    found: dict[str, float] = {}
    for key, pattern in _TARGET_PATTERNS.items():
        m = re.search(pattern, text, flags=re.IGNORECASE)
        if m:
            try:
                found[key] = parse_number(m.group(1))
            except ValueError:
                continue
    return found
