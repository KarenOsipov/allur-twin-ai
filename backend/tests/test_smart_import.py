import io
import json
from datetime import date, datetime, time

from openpyxl import Workbook

from app.domain.plant import ALLUR
from app.ingest import smart
from app.ingest.ai_mapping import parse_reply
from app.services.data_service import CUSTOMER_FILE
from app.services.import_service import ImportService
from tests.conftest import login


def _xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Выпуск"
    ws.append(["Отчёт о выпуске за октябрь"])
    ws.append([])
    ws.append(["Дата смены", "№ смены", "Цех", "План, шт", "Выпуск факт", "Отработано, ч"])
    ws.append([datetime(2026, 10, 1), "1 смена", "Сварочный цех", 120, 117, 7.8])
    ws.append([datetime(2026, 10, 1), "II", "Цех окраски", 120, 110, 7.5])
    ws.append(["Итого", None, None, 240, 227])
    ws.append(["32.10.2026", 1, "Сварка", 120, 100, 7])
    ws2 = wb.create_sheet("Простои")
    ws2.append(["День", "Станок", "Начало", "Причина", "Простой, ч"])
    ws2.append([date(2026, 10, 1), "abb 01", time(9, 15), "Ошибка датчика", 0.5])
    ws2.append([date(2026, 10, 1), "Пресс-77", time(10, 0), "Гидравлика", 1])
    ws3 = wb.create_sheet("Контакты")
    ws3.append(["Телефон", "Комментарий"])
    ws3.append(["+7 777", "звонить утром"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_messy_workbook_is_understood(seeded):
    svc = ImportService(seeded, ALLUR, 300)
    loaded = smart.load("отчёт.xlsx", _xlsx())
    out = svc.to_dict(svc.analyze("отчёт.xlsx", loaded))
    kinds = {b["sheet_name"]: b for b in out["blocks"]}
    prod = kinds["Выпуск"]
    assert prod["kind"] == "production"
    assert out["sheets"][0]["header_row"] == 3
    mapped = {f["name"]: f["column_name"] for f in prod["fields"]}
    assert mapped["area"] == "Цех" and mapped["fact"] == "Выпуск факт" and mapped["shift"] == "№ смены"
    assert prod["counts"]["invalid"] == 1
    assert "32.10.2026" in prod["errors"][0]["message"]
    assert any("итогов" in n for n in out["notes"])
    down = kinds["Простои"]
    assert down["kind"] == "downtime"
    assert any("в часах" in n for n in down["notes"])
    assert down["added"][0]["values"]["Минуты"] == "30"
    assert down["added"][0]["key"].endswith("ABB-01 · 09:15")
    assert down["counts"]["invalid"] == 1
    assert "Контакты" not in kinds


def test_customer_docx_and_resolvers(seeded):
    svc = ImportService(seeded, ALLUR, 300)
    loaded = smart.load(CUSTOMER_FILE.name, CUSTOMER_FILE.read_bytes())
    a = svc.analyze(CUSTOMER_FILE.name, loaded)
    assert [b.kind for b in a.blocks] == ["production", "downtime", "model_plan", "quality"]
    out = svc.to_dict(a)
    assert out["totals"]["new"] == 0 and out["totals"]["update"] == 0
    assert out["targets"]["oee_pct"] == 85
    res = smart.Resolver(ALLUR)
    assert res.area("Цех окраски") == "PAINT"
    assert res.area("ОТК") == "QC"
    assert res.area("Линия сборки №1") == "ASSY"
    assert res.equipment("abb 01") == ("ABB-01", "WELD")
    assert res.model("onix") == "Chevrolet Onix"
    assert smart.parse_shift("Смена II") == 2
    assert smart.parse_month("Октябрь 2026") == "2026-10"
    assert smart.parse_number("1 234,5") == 1234.5


def test_cp1251_csv_with_semicolons():
    raw = "Дата;Участок;Выпущено;Брак\r\n01.10.2026;Окраска;115;4\r\n".encode("cp1251")
    loaded = smart.load("q.csv", raw)
    assert "Windows-1251" in loaded.fmt and "«;»" in loaded.fmt
    smart.detect_header(loaded.sheets[0])
    assert loaded.sheets[0].header == ["Дата", "Участок", "Выпущено", "Брак"]


def test_ai_reply_parsing():
    reply = 'Вот ответ: ```json\n{"kind": "quality", "map": {"day": 0, "area": 1, "produced": 2, "defects": 9}}```'
    assert parse_reply(reply, 4) == {"kind": "quality", "map": {"day": 0, "area": 1, "produced": 2}}
    assert parse_reply("не знаю", 4) is None


def test_analyze_apply_flow(client):
    h = login(client)
    body = "Дата;Смена;Участок;Выпущено;Брак\n01.10.2026;1;Окраска;115;7\n09.10.2026;2;Сборка;118;1\n"
    files = {"file": ("брак.csv", body.encode("cp1251"), "text/csv")}
    res = client.post("/api/v1/data/analyze", headers=h, files=files)
    assert res.status_code == 200, res.text
    a = res.json()
    block = a["blocks"][0]
    assert block["kind"] == "quality"
    assert block["counts"]["update"] == 1 and block["counts"]["new"] == 1
    upd = block["updated"][0]
    assert upd["changes"][0] == {"label": "Брак", "old": "4", "new": "7"}
    assert a["ai"]["used"] is False

    overrides = json.dumps({block["id"]: {"map": {"defects": None}}})
    res = client.post("/api/v1/data/analyze", headers=h, files=files, data={"overrides": overrides, "use_ai": "false"})
    assert res.json()["blocks"][0]["counts"]["new"] == 0

    res = client.post("/api/v1/data/apply", headers=h, files=files)
    assert res.status_code == 200, res.text
    applied = res.json()
    assert applied["new"] == 1 and applied["updated"] == 1
    again = client.post("/api/v1/data/analyze", headers=h, files=files).json()
    assert again["totals"]["same"] == 2 and again["totals"]["new"] == 0
    rows = client.get("/api/v1/data/table/quality?source=import", headers=h).json()
    assert {r["area"] for r in rows} == {"PAINT", "ASSY"}
    summary = client.get("/api/v1/data/summary", headers=h).json()
    assert summary["imports"][0]["filename"] == "брак.csv"


def test_analyze_rejects_and_permissions(client):
    h = login(client)
    director = login(client, "director", "1111")
    files = {"file": ("a.csv", b"x;y\n1;2\n", "text/csv")}
    assert client.post("/api/v1/data/analyze", headers=director, files=files).status_code == 403
    assert client.post("/api/v1/data/apply", headers=director, files=files).status_code == 403
    bad = {"file": ("a.exe", b"MZ", "application/octet-stream")}
    assert client.post("/api/v1/data/analyze", headers=h, files=bad).status_code == 422
    res = client.post("/api/v1/data/analyze", headers=h, files=files).json()
    assert res["blocks"] == []
    text = client.post(
        "/api/v1/data/analyze", headers=h, data={"text": "Модель\tПлан\nJAC J7\t650\n", "use_ai": "false"}
    ).json()
    assert text["blocks"][0]["kind"] == "model_plan"
    assert text["blocks"][0]["updated"][0]["changes"][0]["new"] == "650"
